import os
import re
import subprocess
from docx import Document

def convert_latex_to_docx(latex_file, docx_file):
    """Converte um arquivo LaTeX (.tex) para .docx usando pandoc."""
    try:
        subprocess.run(["pandoc", latex_file, "-o", docx_file], check=True)
    except subprocess.CalledProcessError as e:
        print(f"Erro ao converter {latex_file} para {docx_file}: {e}")


def replace_fields_in_docx(docx_file, field_pattern="\\campo\{(.*?)\}"):
    """Substitui campos específicos no .docx gerado com placeholders."""
    # Compilar o padrão regex para capturar \campo{...}
    regex = re.compile(field_pattern)
    
    # Abrir o arquivo .docx gerado
    doc = Document(docx_file)
    
    for paragraph in doc.paragraphs:
        matches = regex.findall(paragraph.text)
        if matches:
            for match in matches:
                placeholder = f"{{{match}}}"
                paragraph.text = paragraph.text.replace(f"\\campo{{{match}}}", placeholder)
    
    # Salvar o documento com os placeholders inseridos
    doc.save(docx_file)
    print(f"Placeholders inseridos no arquivo: {docx_file}")


def process_files_in_directory(directory):
    """Processa todos os arquivos .tex em um diretório e os converte para .docx."""
    for file_name in os.listdir(directory):
        if file_name.endswith(".tex"):
            latex_file = os.path.join(directory, file_name)
            docx_file = os.path.join(directory, file_name.replace(".tex", ".docx"))
            print(f"Processando {latex_file}...")
            convert_latex_to_docx(latex_file, docx_file)
            replace_fields_in_docx(docx_file)


if __name__ == "__main__":
    directory = input("Digite o caminho do diretório com os arquivos LaTeX: ").strip()
    if os.path.isdir(directory):
        process_files_in_directory(directory)
    else:
        print("Diretório inválido.")
